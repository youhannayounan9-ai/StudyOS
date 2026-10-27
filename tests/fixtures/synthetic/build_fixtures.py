"""Script to build deterministic, non-copyrighted synthetic test document fixtures."""

from pathlib import Path


def build_synthetic_pdf(output_path: Path) -> None:
    """Generate a 2-page deterministic PDF using reportlab."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(output_path), pagesize=letter)

    # Page 1
    c.setFont("Helvetica-Bold", 16)
    c.drawString(72, 720, "Synthetic Lecture: Foundations of Algorithms")
    c.setFont("Helvetica-Bold", 12)
    c.drawString(72, 680, "Section 1: Asymptotic Complexity")
    c.setFont("Helvetica", 10)
    c.drawString(
        72,
        650,
        "Big-O notation describes the upper bound of an algorithm's execution time.",
    )
    c.drawString(
        72,
        630,
        "For example, binary search runs in O(log n) time, while linear search requires O(n).",
    )
    c.drawString(72, 600, "Page 1 of 2: Foundations")
    c.showPage()

    # Page 2
    c.setFont("Helvetica-Bold", 14)
    c.drawString(72, 720, "Section 2: Graph Traversal and Shortest Paths")
    c.setFont("Helvetica", 10)
    c.drawString(
        72,
        680,
        "Breadth-First Search (BFS) utilizes a First-In-First-Out (FIFO) queue.",
    )
    c.drawString(
        72,
        660,
        "Depth-First Search (DFS) relies on recursion or an explicit stack.",
    )
    c.drawString(
        72,
        640,
        "Dijkstra's algorithm computes the shortest path on non-negative weighted graphs.",
    )
    c.drawString(72, 600, "Page 2 of 2: Graph Theory")
    c.showPage()

    c.save()


def build_synthetic_pptx(output_path: Path) -> None:
    """Generate a 4-slide deterministic PPTX presentation with tables using python-pptx."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    blank_layout = prs.slide_layouts[5]  # Title-only layout

    # Slide 1
    slide1 = prs.slides.add_slide(blank_layout)
    slide1.shapes.title.text = "Introduction to Data Structures"
    tx1 = slide1.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(3))
    tf1 = tx1.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "Arrays provide O(1) random access by index."
    p2 = tf1.add_paragraph()
    p2.text = "Linked lists support dynamic resizing with O(1) insertions at the head."

    # Slide 2
    slide2 = prs.slides.add_slide(blank_layout)
    slide2.shapes.title.text = "Tree Architectures and Balancing"
    tx2 = slide2.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(3))
    tf2 = tx2.text_frame
    p3 = tf2.paragraphs[0]
    p3.text = "Binary Search Trees (BST) maintain sorted order."
    p4 = tf2.add_paragraph()
    p4.text = "Balanced trees like AVL and Red-Black trees guarantee O(log n) search depth."

    # Slide 3
    slide3 = prs.slides.add_slide(blank_layout)
    slide3.shapes.title.text = "Hash Tables and Collision Resolution"
    tx3 = slide3.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(3))
    tf3 = tx3.text_frame
    p5 = tf3.paragraphs[0]
    p5.text = "Hash maps resolve collisions via chaining or open addressing with linear probing."

    # Slide 4: Table with algorithm complexity
    slide4 = prs.slides.add_slide(blank_layout)
    slide4.shapes.title.text = "Algorithm Complexity Summary"
    rows, cols = 3, 3
    left = Inches(1)
    top = Inches(2)
    width = Inches(8)
    height = Inches(2)
    table_shape = slide4.shapes.add_table(rows, cols, left, top, width, height)
    tbl = table_shape.table
    tbl.cell(0, 0).text = "Algorithm"
    tbl.cell(0, 1).text = "Average Time"
    tbl.cell(0, 2).text = "Worst Space"
    tbl.cell(1, 0).text = "QuickSort"
    tbl.cell(1, 1).text = "O(n log n)"
    tbl.cell(1, 2).text = "O(log n)"
    tbl.cell(2, 0).text = "MergeSort"
    tbl.cell(2, 1).text = "O(n log n)"
    tbl.cell(2, 2).text = "O(n)"

    prs.save(str(output_path))


def build_synthetic_docx(output_path: Path) -> None:
    """Generate a multi-section DOCX with headings and tables using python-docx."""
    import docx

    doc = docx.Document()
    doc.add_heading("Relational Database Internals", level=1)
    doc.add_paragraph(
        "Relational engines enforce ACID properties: Atomicity, Consistency, Isolation, and Durability."
    )

    doc.add_heading("Indexing Structures and Storage Layout", level=2)
    doc.add_paragraph(
        "B-Trees and B+ Trees reduce disk I/O by maximizing fan-out across disk blocks."
    )

    table = doc.add_table(rows=3, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Index Type"
    hdr_cells[1].text = "Time Complexity"
    hdr_cells[2].text = "Storage Layout"

    row1_cells = table.rows[1].cells
    row1_cells[0].text = "B+ Tree"
    row1_cells[1].text = "O(log n)"
    row1_cells[2].text = "Block-oriented tree"

    row2_cells = table.rows[2].cells
    row2_cells[0].text = "Hash Index"
    row2_cells[1].text = "O(1)"
    row2_cells[2].text = "Bucket array"

    doc.save(str(output_path))


def build_synthetic_txt(output_path: Path) -> None:
    """Generate a clean UTF-8 text document with markdown headings and unicode math."""
    content = """# Computer Systems: Memory Hierarchy

Modern computer architectures employ multi-level cache hierarchies: L1, L2, and L3 caches.
Cache misses incur main memory latency (approximately 50–100 ns).

## Virtual Memory and Paging

Virtual memory provides process isolation using page tables.
Translation Lookaside Buffers (TLB) accelerate virtual-to-physical address mapping.
Mathematical latency model: Δt = α · tc + (1 - α) · tm.
"""
    output_path.write_text(content, encoding="utf-8")


def generate_all_fixtures(target_dir: Path | None = None) -> dict[str, Path]:
    """Generate all synthetic test fixtures into the specified directory."""
    dir_path = target_dir or Path(__file__).parent
    dir_path.mkdir(parents=True, exist_ok=True)

    pdf_path = dir_path / "synthetic_lecture.pdf"
    pptx_path = dir_path / "synthetic_slides.pptx"
    docx_path = dir_path / "synthetic_notes.docx"
    txt_path = dir_path / "synthetic_doc.txt"

    build_synthetic_pdf(pdf_path)
    build_synthetic_pptx(pptx_path)
    build_synthetic_docx(docx_path)
    build_synthetic_txt(txt_path)

    return {
        "pdf": pdf_path,
        "pptx": pptx_path,
        "docx": docx_path,
        "txt": txt_path,
    }


if __name__ == "__main__":
    fixtures = generate_all_fixtures()
    for name, p in fixtures.items():
        print(f"Generated {name.upper()} fixture: {p.name} ({p.stat().st_size} bytes)")
