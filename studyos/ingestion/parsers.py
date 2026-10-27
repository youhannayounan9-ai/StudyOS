"""Format-specific document parsers for PDF, DOCX, PPTX, and TXT files."""

from __future__ import annotations

from abc import ABC, abstractmethod
import hashlib
from pathlib import Path
import re
from typing import Any

from studyos.ingestion.models import (
    CorruptFileError,
    EmptyFileError,
    EncryptedFileError,
    ExtractionEmptyError,
    FileTooLargeError,
    PageElement,
    ParsedDocument,
    UnsupportedFileTypeError,
)


class DocumentParser(ABC):
    """Abstract base class for format-specific document parsers."""

    def __init__(self, max_file_size_bytes: int = 100 * 1024 * 1024) -> None:
        self.max_file_size_bytes = max_file_size_bytes

    def validate_file(self, file_path: Path) -> tuple[int, str]:
        """Validate file existence, size, and compute SHA-256 hash."""
        if not file_path.is_file():
            raise FileNotFoundError(f"File not found: {file_path}")

        file_size = file_path.stat().st_size
        if file_size == 0:
            raise EmptyFileError(f"File is empty (0 bytes): {file_path.name}")

        if file_size > self.max_file_size_bytes:
            raise FileTooLargeError(
                f"File size ({file_size} bytes) exceeds limit of {self.max_file_size_bytes} bytes: {file_path.name}"
            )

        sha256 = hashlib.sha256()
        with open(file_path, "rb") as f:
            while chunk := f.read(65536):
                sha256.update(chunk)

        return file_size, sha256.hexdigest()

    @abstractmethod
    def parse(self, file_path: Path, filename: str | None = None) -> ParsedDocument:
        """Parse file into a normalized ParsedDocument entity."""
        pass


class PDFParser(DocumentParser):
    """Parser for Adobe Portable Document Format (.pdf) files using pypdf."""

    def parse(self, file_path: Path, filename: str | None = None) -> ParsedDocument:
        file_size, file_hash = self.validate_file(file_path)
        actual_filename = filename or file_path.name

        try:
            import pypdf
        except ImportError as e:
            raise RuntimeError("pypdf is required to parse PDF files.") from e

        try:
            reader = pypdf.PdfReader(str(file_path))
        except Exception as e:
            raise CorruptFileError(f"Failed to read PDF file '{actual_filename}': {e}") from e

        if reader.is_encrypted:
            try:
                # Attempt empty password decryption
                decrypted = reader.decrypt("")
                if decrypted == pypdf.PasswordType.NOT_DECRYPTED:
                    raise EncryptedFileError(
                        f"PDF file '{actual_filename}' is password-protected and cannot be parsed."
                    )
            except Exception as e:
                if isinstance(e, EncryptedFileError):
                    raise
                raise EncryptedFileError(
                    f"PDF file '{actual_filename}' is encrypted and password was not provided."
                ) from e

        elements: list[PageElement] = []
        total_chars = 0
        total_pages = len(reader.pages)

        if total_pages == 0:
            raise ExtractionEmptyError(f"PDF file '{actual_filename}' contains 0 pages.")

        for idx, page in enumerate(reader.pages, start=1):
            try:
                page_text = page.extract_text() or ""
            except Exception as e:
                # Corrupt page stream
                elements.append(
                    PageElement(
                        text="",
                        location_type="page",
                        location_index=idx,
                        warnings=[f"Page extraction error on page {idx}: {e}"],
                    )
                )
                continue

            cleaned = page_text.strip()
            total_chars += len(cleaned)
            elements.append(
                PageElement(
                    text=cleaned,
                    location_type="page",
                    location_index=idx,
                )
            )

        if total_chars == 0:
            raise ExtractionEmptyError(
                f"PDF file '{actual_filename}' contains no extractable text across {total_pages} pages. "
                "Scanned documents requiring OCR are not supported in this release."
            )

        warnings: list[str] = []
        avg_chars_per_page = total_chars / total_pages
        if avg_chars_per_page < 50.0:
            warnings.append(
                f"ScannedDocumentWarning: Average extracted page text is below 50 characters "
                f"({avg_chars_per_page:.1f} chars/page). Document may contain scanned bitmap images or complex layouts."
            )

        return ParsedDocument(
            filename=actual_filename,
            file_type="pdf",
            sha256_hash=file_hash,
            file_size_bytes=file_size,
            elements=elements,
            warnings=warnings,
        )


class DOCXParser(DocumentParser):
    """Parser for Microsoft Word (.docx) documents using python-docx."""

    def parse(self, file_path: Path, filename: str | None = None) -> ParsedDocument:
        file_size, file_hash = self.validate_file(file_path)
        actual_filename = filename or file_path.name

        try:
            import docx
        except ImportError as e:
            raise RuntimeError("python-docx is required to parse DOCX files.") from e

        try:
            doc = docx.Document(str(file_path))
        except Exception as e:
            raise CorruptFileError(f"Failed to read DOCX file '{actual_filename}': {e}") from e

        elements: list[PageElement] = []
        current_heading: str | None = None
        current_text_parts: list[str] = []
        section_idx = 1

        def flush_section() -> None:
            nonlocal section_idx, current_text_parts
            text = "\n\n".join(t for t in current_text_parts if t.strip()).strip()
            if text:
                elements.append(
                    PageElement(
                        text=text,
                        location_type="section",
                        location_index=section_idx,
                        location_title=current_heading,
                    )
                )
                section_idx += 1
            current_text_parts = []

        # Iterate through block elements in body order (paragraphs and tables)
        for element in doc.element.body:
            tag = element.tag.split("}")[-1] if "}" in element.tag else element.tag

            if tag == "p":
                # Paragraph
                p = docx.text.paragraph.Paragraph(element, doc)
                p_text = p.text.strip()
                if not p_text:
                    continue

                style_name = p.style.name if p.style else ""
                if "heading" in style_name.lower():
                    # Heading boundary -> flush preceding section
                    flush_section()
                    current_heading = p_text
                    current_text_parts.append(p_text)
                else:
                    current_text_parts.append(p_text)

            elif tag == "tbl":
                # Table
                table = docx.table.Table(element, doc)
                table_lines: list[str] = []
                for row in table.rows:
                    row_cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                    if any(row_cells):
                        table_lines.append("| " + " | ".join(row_cells) + " |")
                if table_lines:
                    current_text_parts.append("\n".join(table_lines))

        flush_section()

        if not elements or sum(len(el.text) for el in elements) == 0:
            raise ExtractionEmptyError(f"DOCX file '{actual_filename}' contains no extractable text.")

        return ParsedDocument(
            filename=actual_filename,
            file_type="docx",
            sha256_hash=file_hash,
            file_size_bytes=file_size,
            elements=elements,
        )


class PPTXParser(DocumentParser):
    """Parser for Microsoft PowerPoint (.pptx) presentations using python-pptx."""

    def parse(self, file_path: Path, filename: str | None = None) -> ParsedDocument:
        file_size, file_hash = self.validate_file(file_path)
        actual_filename = filename or file_path.name

        try:
            import pptx
        except ImportError as e:
            raise RuntimeError("python-pptx is required to parse PPTX files.") from e

        try:
            prs = pptx.Presentation(str(file_path))
        except Exception as e:
            raise CorruptFileError(f"Failed to read PPTX file '{actual_filename}': {e}") from e

        elements: list[PageElement] = []
        total_slides = len(prs.slides)

        if total_slides == 0:
            raise ExtractionEmptyError(f"PPTX file '{actual_filename}' contains 0 slides.")

        for idx, slide in enumerate(prs.slides, start=1):
            slide_title: str | None = None
            text_parts: list[str] = []

            # Check for title shape
            if slide.shapes.title and slide.shapes.title.has_text_frame:
                title_text = slide.shapes.title.text.strip()
                if title_text:
                    slide_title = title_text

            # Iterate all shapes on the slide (tables and text frames)
            for shape in slide.shapes:
                if shape.has_table:
                    table_lines: list[str] = []
                    for row in shape.table.rows:
                        row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                        if any(row_cells):
                            table_lines.append("| " + " | ".join(row_cells) + " |")
                    if table_lines:
                        text_parts.append("\n".join(table_lines))
                elif shape.has_text_frame:
                    shape_text = shape.text.strip()
                    if shape_text:
                        # Don't duplicate slide title if it's already captured
                        if slide_title and shape_text == slide_title:
                            continue
                        text_parts.append(shape_text)

            combined_text = "\n\n".join(text_parts).strip()
            # If slide has title but no other body text, record title as text
            if not combined_text and slide_title:
                combined_text = slide_title

            elements.append(
                PageElement(
                    text=combined_text,
                    location_type="slide",
                    location_index=idx,
                    location_title=slide_title,
                )
            )

        total_chars = sum(len(el.text) for el in elements)
        if total_chars == 0:
            raise ExtractionEmptyError(
                f"PPTX file '{actual_filename}' contains no extractable text across {total_slides} slides."
            )

        return ParsedDocument(
            filename=actual_filename,
            file_type="pptx",
            sha256_hash=file_hash,
            file_size_bytes=file_size,
            elements=elements,
        )


class TXTParser(DocumentParser):
    """Parser for plain text (.txt) and Markdown files with UTF-8 encoding."""

    def parse(self, file_path: Path, filename: str | None = None) -> ParsedDocument:
        file_size, file_hash = self.validate_file(file_path)
        actual_filename = filename or file_path.name

        try:
            # Handle UTF-8 with or without BOM (utf-8-sig automatically strips BOM)
            with open(file_path, "r", encoding="utf-8-sig") as f:
                content = f.read()
        except UnicodeDecodeError as e:
            raise CorruptFileError(
                f"File '{actual_filename}' is not valid UTF-8 encoded text: {e}"
            ) from e
        except Exception as e:
            raise CorruptFileError(f"Failed to read TXT file '{actual_filename}': {e}") from e

        cleaned = content.strip()
        if not cleaned:
            raise ExtractionEmptyError(f"TXT file '{actual_filename}' contains no text.")

        # Split by markdown headers (# Header) into sections if present
        lines = cleaned.splitlines()
        elements: list[PageElement] = []
        current_heading: str | None = None
        current_lines: list[str] = []
        sec_idx = 1

        def flush_section() -> None:
            nonlocal sec_idx, current_lines
            sec_text = "\n".join(current_lines).strip()
            if sec_text:
                elements.append(
                    PageElement(
                        text=sec_text,
                        location_type="section",
                        location_index=sec_idx,
                        location_title=current_heading,
                    )
                )
                sec_idx += 1
            current_lines = []

        header_pattern = re.compile(r"^(#{1,6})\s+(.+)$")
        has_headers = any(header_pattern.match(l.strip()) for l in lines)

        if has_headers:
            for line in lines:
                match = header_pattern.match(line.strip())
                if match:
                    flush_section()
                    current_heading = match.group(2).strip()
                    current_lines.append(line)
                else:
                    current_lines.append(line)
            flush_section()
        else:
            # No markdown headers -> single section
            elements.append(
                PageElement(
                    text=cleaned,
                    location_type="section",
                    location_index=1,
                    location_title=None,
                )
            )

        return ParsedDocument(
            filename=actual_filename,
            file_type="txt",
            sha256_hash=file_hash,
            file_size_bytes=file_size,
            elements=elements,
        )


def get_parser_for_file(
    file_path: Path, max_file_size_bytes: int = 100 * 1024 * 1024
) -> DocumentParser:
    """Factory returning the appropriate DocumentParser based on file extension."""
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return PDFParser(max_file_size_bytes=max_file_size_bytes)
    elif suffix in (".docx", ".doc"):
        if suffix == ".doc":
            raise UnsupportedFileTypeError(
                "Legacy .doc binary format is not supported. Please convert to .docx."
            )
        return DOCXParser(max_file_size_bytes=max_file_size_bytes)
    elif suffix in (".pptx", ".ppt"):
        if suffix == ".ppt":
            raise UnsupportedFileTypeError(
                "Legacy .ppt binary format is not supported. Please convert to .pptx."
            )
        return PPTXParser(max_file_size_bytes=max_file_size_bytes)
    elif suffix in (".txt", ".md"):
        return TXTParser(max_file_size_bytes=max_file_size_bytes)
    else:
        raise UnsupportedFileTypeError(
            f"Unsupported file format '{suffix}'. Supported formats are: .pdf, .docx, .pptx, .txt, .md."
        )
